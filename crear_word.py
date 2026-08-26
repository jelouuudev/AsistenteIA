from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

doc.add_paragraph('LABORATORIO 01 - BIG DATA')
doc.add_paragraph('ALUMNO: [TU NOMBRE AQUI]')
doc.add_paragraph('')
doc.add_paragraph('ACTIVIDAD DE COMPROBACION:')

# Punto 6
preguntas = [
    ("\u00bfCu\u00e1ntos registros contiene el dataset?",
     "Tiene 150 registros (filas). Se observa en df.shape como (150, 7) y en RangeIndex: 150 entries."),
    ("\u00bfCu\u00e1ntas columnas contiene?",
     "Contiene 7 columnas."),
    ("\u00bfCu\u00e1les son los nombres de las columnas?",
     "ID, Edad, Genero_Preferido, Pelicula_Preferida, Frecuencia_Cine, Formato_Preferido y Calificacion_Cine."),
    ("\u00bfQu\u00e9 tipo de dato tiene la columna Cantidad?",
     "Este dataset no contiene la columna Cantidad. Las columnas num\u00e9ricas son: ID (int64), Edad (int64) y Calificacion_Cine (int64)."),
    ("\u00bfQu\u00e9 tipo de dato tiene inicialmente la columna Fecha?",
     "Este dataset no contiene la columna Fecha. Las columnas de tipo texto (object) son: Genero_Preferido, Pelicula_Preferida, Frecuencia_Cine y Formato_Preferido."),
    ("Seg\u00fan describe(), \u00bfcu\u00e1l es el precio m\u00e1ximo registrado?",
     "Este dataset no contiene la columna Precio. La columna Edad m\u00e1xima es [EJECUTAR EN COLAB]. La columna Calificacion_Cine m\u00e1xima es [EJECUTAR EN COLAB]."),
    ("Seg\u00fan describe(), \u00bfcu\u00e1l es la cantidad m\u00e1xima registrada en una venta?",
     "Este dataset no contiene la columna Cantidad. La Edad m\u00e1xima registrada es [EJECUTAR EN COLAB]."),
]

for pregunta, respuesta in preguntas:
    p = doc.add_paragraph(pregunta, style='List Paragraph')
    doc.add_paragraph(respuesta)

doc.add_paragraph('')
doc.add_paragraph('')

# Punto 7
doc.add_paragraph('RESULTADO ESPERADO:')
doc.add_paragraph('')

doc.add_paragraph('Cargar el archivo CSV en Google Colab y crear un DataFrame utilizando Pandas.', style='List Paragraph')
doc.add_paragraph('[PEGAR CAPTURA DE df.head() AQUI]', style='List Paragraph')
doc.add_paragraph('')

doc.add_paragraph('Visualizar una muestra de los registros.', style='List Paragraph')
doc.add_paragraph('[PEGAR CAPTURA DE df.head() AQUI]', style='List Paragraph')
doc.add_paragraph('')

doc.add_paragraph('Identificar el n\u00famero de filas y columnas del dataset.', style='List Paragraph')
doc.add_paragraph('Tiene 150 registros (filas) y 7 columnas.', style='List Paragraph')
doc.add_paragraph('[PEGAR CAPTURA DE df.shape AQUI]', style='List Paragraph')
doc.add_paragraph('')

doc.add_paragraph('Reconocer los nombres y tipos de datos de las columnas.', style='List Paragraph')
doc.add_paragraph('[PEGAR CAPTURA DE df.info() AQUI]', style='List Paragraph')
doc.add_paragraph('Las columnas ID, Edad y Calificacion_Cine son de tipo entero (int64). Las columnas Genero_Preferido, Pelicula_Preferida, Frecuencia_Cine y Formato_Preferido son de tipo texto (object).')
doc.add_paragraph('')
doc.add_paragraph('')

doc.add_paragraph('Obtener estad\u00edsticas descriptivas b\u00e1sicas de las variables num\u00e9ricas.', style='List Paragraph')
doc.add_paragraph('[PEGAR CAPTURA DE df.describe() AQUI]', style='List Paragraph')
doc.add_paragraph('')
doc.add_paragraph('Interpretaci\u00f3n: Seg\u00fan df.describe(), la muestra analiza a 150 personas con un promedio de edad de [PONER VALOR] a\u00f1os (rango entre [MIN] y [MAX] a\u00f1os). La calificaci\u00f3n promedio es [PONER VALOR] sobre 5.')
doc.add_paragraph('')

doc.save('LAB01_Resuestas.docx')
print("Documento creado")
